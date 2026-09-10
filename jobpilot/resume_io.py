"""Read common resume formats and export a simple ATS-friendly DOCX."""

from __future__ import annotations

from pathlib import Path


def read_resume(path: str | Path) -> str:
    """Extract text from PDF or DOCX resumes without changing their facts."""
    source = Path(path).expanduser()
    if not source.is_file():
        raise FileNotFoundError(source)
    suffix = source.suffix.casefold()
    if suffix == ".pdf":
        from pypdf import PdfReader

        return "\n\n".join(page.extract_text() or "" for page in PdfReader(source).pages).strip()
    if suffix == ".docx":
        from docx import Document

        document = Document(source)
        return "\n".join(paragraph.text for paragraph in document.paragraphs).strip()
    if suffix in {".txt", ".md"}:
        return source.read_text(encoding="utf-8").strip()
    raise ValueError("Unsupported resume format. Use PDF, DOCX, TXT, or Markdown.")


def write_ats_docx(markdown_text: str, output_path: str | Path) -> Path:
    """Write plain, single-column DOCX output suitable for ATS parsing."""
    from docx import Document
    from docx.shared import Pt

    if not markdown_text.strip():
        raise ValueError("markdown_text must not be empty")
    output = Path(output_path).expanduser()
    output.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    style = document.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(10.5)
    for raw_line in markdown_text.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        if line.startswith("#"):
            paragraph = document.add_paragraph()
            paragraph.add_run(line.lstrip("# ")).bold = True
        elif line.startswith("- "):
            document.add_paragraph(line[2:], style="List Bullet")
        else:
            document.add_paragraph(line)
    document.save(output)
    return output
