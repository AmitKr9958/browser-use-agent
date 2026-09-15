"""Resume extraction regression tests."""

from docx import Document

from browser_agent.jobpilot.profile import extract_resume_text, infer_contact_profile


def test_docx_table_content_is_included(tmp_path) -> None:
    path = tmp_path / "resume.docx"
    document = Document()
    document.add_paragraph("Candidate Name")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Employer"
    table.cell(0, 1).text = "Example Technologies"
    document.save(path)
    text = extract_resume_text(path)
    assert "Example Technologies" in text


def test_contact_inference_remains_conservative() -> None:
    profile = infer_contact_profile("Candidate Name\nname@example.com\n+91 98765 43210")
    assert profile.name == "Candidate Name"
    assert profile.email == "name@example.com"
    assert profile.phone == "+91 98765 43210"
    assert profile.work_authorization == ""
