"""Resume tailoring and cover-letter generation for JobPilot."""

from __future__ import annotations

import re
from pathlib import Path

from browser_use import ChatGoogle
from browser_use.llm.messages import UserMessage
from docx import Document

from .evidence import format_resume_evidence
from .models import ContactProfile, JobDescription

_EMAIL_RE = re.compile(r"(?<![\w.+-])[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}(?![\w.-])", re.I)
_PHONE_RE = re.compile(r"(?<!\d)(?:\+?\d[\d ()-]{8,}\d)(?!\d)")
_URL_RE = re.compile(r"https?://[^\s)]+", re.I)


def tailor_resume_text(resume_text: str, missing_keywords: tuple[str, ...], *, max_additions: int = 15) -> str:
    """Create a deterministic ATS review draft without inventing experience."""
    if not resume_text.strip():
        raise ValueError("resume_text must not be empty")
    if max_additions < 0:
        raise ValueError("max_additions must be non-negative")
    additions = tuple(dict.fromkeys(word.strip() for word in missing_keywords if word.strip()))[:max_additions]
    if not additions:
        return resume_text.strip()
    review = "\n".join(f"- {word}" for word in additions)
    return (
        f"{resume_text.strip()}\n\n"
        "[JOBPILOT REVIEW — add only if accurate]\n"
        "Relevant keywords from the job description not found in the current resume:\n"
        f"{review}"
    )


def validate_generated_resume(source_resume: str, generated_resume: str) -> str:
    """Reject unsafe LLM resume output containing newly invented high-risk identifiers."""
    source = source_resume.strip()
    generated = generated_resume.strip()
    if not source or not generated:
        raise ValueError("source and generated resumes must not be empty")

    def normalize(values: list[str]) -> set[str]:
        return {value.rstrip(".,;:)").casefold() for value in values}

    if not normalize(_EMAIL_RE.findall(generated)).issubset(normalize(_EMAIL_RE.findall(source))):
        return source
    if not normalize(_PHONE_RE.findall(generated)).issubset(normalize(_PHONE_RE.findall(source))):
        return source
    if not normalize(_URL_RE.findall(generated)).issubset(normalize(_URL_RE.findall(source))):
        return source
    if len(generated) > max(len(source) * 3, 20000):
        return source
    return generated


async def tailor_resume_with_llm(
    resume_text: str,
    job_description: str,
    *,
    model: str = "gemini-3.6-flash",
) -> str:
    """Generate an ATS-focused resume draft while explicitly prohibiting invented facts."""
    if not resume_text.strip() or not job_description.strip():
        raise ValueError("resume_text and job_description must not be empty")
    evidence = format_resume_evidence(resume_text)
    prompt = f"""
Rewrite the resume below for the target job description.

STRICT RULES:
- Treat SOURCE-BACKED EVIDENCE as the authoritative fact boundary.
- Use only facts, employers, titles, dates, education, certifications, skills, and achievements present in the source resume/evidence.
- Do not invent metrics, responsibilities, technologies, employers, credentials, locations, or years of experience.
- You may reorder sections, tighten wording, remove repetition, and emphasize experience that is already present.
- If a job keyword is not supported by the source evidence, do not add it as a claimed skill.
- Never change, invent, or add an email address, phone number, URL, employer, credential, date, metric, or identity detail.
- Keep the output ATS-friendly: plain text headings, concise bullets, no tables, no icons, no graphics.
- Return only the revised resume text, with no explanation.

TARGET JOB:
{job_description}

SOURCE-BACKED EVIDENCE:
{evidence}

SOURCE RESUME:
{resume_text}
""".strip()
    llm = ChatGoogle(model=model)
    response = await llm.ainvoke([UserMessage(content=prompt)])
    output = str(response.completion).strip()
    if not output:
        raise RuntimeError("resume tailoring model returned empty output")
    return validate_generated_resume(resume_text, output)


def write_resume_docx(resume_text: str, output_path: str | Path) -> Path:
    """Write a clean, ATS-friendly DOCX draft without tables or graphics."""
    if not resume_text.strip():
        raise ValueError("resume_text must not be empty")
    path = Path(output_path).expanduser()
    path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    for raw_line in resume_text.strip().splitlines():
        line = raw_line.strip()
        if not line:
            continue
        if line.startswith("### "):
            document.add_heading(line[4:].strip(), level=2)
        elif line.startswith("## "):
            document.add_heading(line[3:].strip(), level=1)
        elif line.startswith("# "):
            document.add_heading(line[2:].strip(), level=1)
        elif line.startswith("- ") or line.startswith("* "):
            document.add_paragraph(line[2:].strip(), style="List Bullet")
        else:
            document.add_paragraph(line)
    document.save(str(path))
    return path


def build_cover_letter(job: JobDescription, profile: ContactProfile, *, resume_text: str = "") -> str:
    """Build a concise cover letter from supplied facts only."""
    name = profile.name or "Applicant"
    evidence = "I would welcome the opportunity to discuss how my experience can contribute to this role."
    if resume_text.strip():
        evidence = "My background includes experience relevant to the requirements described in the posting, and I would welcome the opportunity to discuss the strongest examples."
    return (
        "Dear Hiring Team,\n\n"
        f"I am interested in the {job.title} opportunity at {job.company}. "
        f"{evidence}\n\n"
        "I am particularly interested in the scope described in the job posting and the opportunity to contribute to the team. "
        "Please find my resume attached for consideration.\n\n"
        "Thank you for your time and consideration.\n\n"
        f"Regards,\n{name}"
    )


async def build_cover_letter_with_llm(
    job: JobDescription,
    profile: ContactProfile,
    resume_text: str,
    *,
    model: str = "gemini-3.6-flash",
) -> str:
    """Generate a targeted cover-letter draft using only supplied resume/profile facts."""
    if not resume_text.strip():
        return build_cover_letter(job, profile)
    evidence = format_resume_evidence(resume_text)
    prompt = f"""
Write a concise cover letter for this job using only facts present in the supplied profile, resume, and SOURCE-BACKED EVIDENCE.
Do not invent employers, achievements, metrics, skills, credentials, dates, or experience.
Never invent or change contact details.
Keep it professional, specific to the role, ATS-friendly, and under 250 words.
Return only the cover letter.

JOB:
Title: {job.title}
Company: {job.company}
Description:
{job.description}

PROFILE:
Name: {profile.name}
Location: {profile.location}

SOURCE-BACKED EVIDENCE:
{evidence}

RESUME:
{resume_text}
""".strip()
    llm = ChatGoogle(model=model)
    response = await llm.ainvoke([UserMessage(content=prompt)])
    output = str(response.completion).strip()
    if not output:
        raise RuntimeError("cover letter model returned empty output")
    return output
