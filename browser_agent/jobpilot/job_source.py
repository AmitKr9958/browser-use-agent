"""Job posting ingestion helpers for JobPilot."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any, cast

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from .models import JobDescription


@dataclass(frozen=True)
class JobPostingRow:
    """Normalized workbook row, preserving source metadata."""

    row_number: int
    job: JobDescription
    match_score: float | None = None
    key_matching_skills: str = ""
    date_found: str = ""
    job_posting_date: str = ""


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value).strip()


def _score(value: Any) -> float | None:
    if value in (None, ""):
        return None
    try:
        score = float(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"invalid Match Score: {value!r}") from exc
    if not 0 <= score <= 1:
        raise ValueError(f"Match Score must be between 0 and 1: {value!r}")
    return score


def load_job_postings_xlsx(path: str | Path) -> list[JobPostingRow]:
    """Load the supported JobPilot workbook format with strict headers."""
    workbook_path = Path(path).expanduser()
    if not workbook_path.is_file():
        raise FileNotFoundError(f"job posting workbook not found: {workbook_path}")
    if workbook_path.suffix.lower() != ".xlsx":
        raise ValueError("job posting source must be an .xlsx workbook")

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        sheet = cast(Worksheet, workbook.active)
        rows = sheet.iter_rows(values_only=True)
        try:
            headers = [_text(value) for value in next(rows)]
        except StopIteration as exc:
            raise ValueError("job posting workbook is empty") from exc
        index = {header: position for position, header in enumerate(headers) if header}
        required = {"Company Name", "Job Title", "Location", "Job Posting URL"}
        missing = sorted(required - index.keys())
        if missing:
            raise ValueError(f"job posting workbook missing required headers: {', '.join(missing)}")

        postings: list[JobPostingRow] = []
        for row_number, values in enumerate(rows, start=2):
            get = lambda header: _text(values[index[header]]) if index[header] < len(values) else ""
            company = get("Company Name")
            title = get("Job Title")
            url = get("Job Posting URL")
            if not any((company, title, url)):
                continue
            if not title or not company or not url:
                raise ValueError(f"row {row_number}: Company Name, Job Title and Job Posting URL are required")
            skills = get("Key Matching Skills") if "Key Matching Skills" in index else ""
            postings.append(JobPostingRow(row_number=row_number, job=JobDescription(title=title, company=company, description=skills, url=url, location=get("Location") if "Location" in index else ""), match_score=_score(values[index["Match Score"]]) if "Match Score" in index and index["Match Score"] < len(values) else None, key_matching_skills=skills, date_found=get("Date Found") if "Date Found" in index else "", job_posting_date=get("Job Posting Date") if "Job Posting Date" in index else ""))
        return postings
    finally:
        workbook.close()


def get_job_posting_xlsx(path: str | Path, row_number: int) -> JobPostingRow:
    """Return a 1-based Excel row (including the header row)."""
    if row_number < 2:
        raise ValueError("row_number must be 2 or greater")
    for posting in load_job_postings_xlsx(path):
        if posting.row_number == row_number:
            return posting
    raise IndexError(f"job posting row not found: {row_number}")
